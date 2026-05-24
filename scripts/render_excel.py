"""Render analysis results as an Excel workbook."""

import io
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path


HEADER_FILL = PatternFill(start_color='2563eb', end_color='2563eb', fill_type='solid')
HEADER_FONT = Font(color='ffffff', bold=True, size=11)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)


def _write_sheet(ws, headers: List[str], rows: List[list], title: Optional[str] = None):
    start_row = 1
    if title:
        ws['A1'] = title
        ws['A1'].font = Font(bold=True, size=13)
        start_row = 3

    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=j, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
        cell.border = THIN_BORDER

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = ws.cell(row=start_row + 1 + i, column=j + 1, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal='center')
            # Color code change columns
            if isinstance(val, str) and val.startswith('+'):
                cell.font = Font(color='16a34a')
            elif isinstance(val, str) and val.startswith('-') and val != '-':
                cell.font = Font(color='dc2626')

    # Auto-width
    for col in ws.columns:
        max_len = 0
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 25)


def _build_workbook(result: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    m = result['meta']

    # 1. Funnel overview
    ws = wb.create_sheet('漏斗总览')
    headers = ['指标', '本周', '前4周均值', '变化']
    rows = [[r['metric'], r['this_week'], r['prev_4weeks'], r['change']] for r in result['funnel']]
    _write_sheet(ws, headers, rows, f'漏斗总览 | {m["report_start"]} ~ {m["report_end"]}')

    # 2. Per capita
    ws = wb.create_sheet('人均金额')
    headers = ['指标', '本周', '前4周均值', '变化']
    rows = [[r['metric'], r['this_week'], r['prev_4weeks'], r['change']] for r in result['per_capita']]
    _write_sheet(ws, headers, rows)

    # 3. Spot vs futures
    ws = wb.create_sheet('现货vs合约')
    sf = result['spot_futures']
    headers = ['类型', '交易人数', '交易金额', '占比']
    rows = [
        ['现货(本周)', sf['this_week']['spot_count'], round(sf['this_week']['spot_amount'], 2), f'{sf["this_week"]["spot_pct"]:.1%}'],
        ['合约(本周)', sf['this_week']['futures_count'], round(sf['this_week']['futures_amount'], 2), f'{sf["this_week"]["futures_pct"]:.1%}'],
        ['现货(前4周)', sf['prev_4weeks']['spot_count'], round(sf['prev_4weeks']['spot_amount'], 2), f'{sf["prev_4weeks"]["spot_pct"]:.1%}'],
        ['合约(前4周)', sf['prev_4weeks']['futures_count'], round(sf['prev_4weeks']['futures_amount'], 2), f'{sf["prev_4weeks"]["futures_pct"]:.1%}'],
    ]
    _write_sheet(ws, headers, rows)

    # 4. Channel ranking
    ws = wb.create_sheet('渠道排名')
    headers = ['排名', '渠道', '注册量', '充值率', '交易率', '人均交易', '交易金额', '排名变化']
    rows = [[r['rank'], r['channel'], r['reg'], r['deposit_rate'], r['trade_rate'],
             r['avg_trade'], r['trade_amount'], r['rank_change']] for r in result['channels']['rows']]
    _write_sheet(ws, headers, rows)

    # 5. Country top 10
    ws = wb.create_sheet('国家Top10')
    headers = ['排名', '国家', '注册量', '充值率', '交易率', '人均充值', '人均交易', '充值率变化', '交易率变化']
    rows = [[r['rank'], r['country'], r['reg'], r['deposit_rate'], r['trade_rate'],
             r['avg_deposit'], r['avg_trade'], r['dep_change'], r['trade_change']] for r in result['countries']]
    _write_sheet(ws, headers, rows)

    # 6. Wool analysis
    ws = wb.create_sheet('羊毛分析')
    headers = ['渠道', '全部注册', '去羊毛注册', '羊毛量', '羊毛比例', '变化']
    rows = [[r['channel'], r['total'], r['filtered'], r['wool'], r['wool_pct'], r['change']]
            for r in result['wool']['channels']]
    ov = result['wool']['overall']
    rows.append(['总计', ov['tw_total'], ov['tw_total'] - ov['tw_wool'], ov['tw_wool'], ov['tw_pct'], ''])
    _write_sheet(ws, headers, rows)

    # 7. Monthly trend
    ws = wb.create_sheet('月度趋势')
    monthly = result['monthly']
    headers = ['月份', '注册人数', 'KYC1率', 'KYC2率', '充值率', '交易率', '充值金额', '交易金额']
    rows = []
    for i, month in enumerate(monthly['months']):
        rows.append([
            month,
            int(monthly['reg'][i]),
            f'{monthly["kyc1_rate"][i]:.2%}',
            f'{monthly["kyc2_rate"][i]:.2%}',
            f'{monthly["deposit_rate"][i]:.2%}',
            f'{monthly["trade_rate"][i]:.2%}',
            round(monthly['deposit_amount'][i], 2),
            round(monthly['trade_amount'][i], 2),
        ])
    _write_sheet(ws, headers, rows)

    return wb


def render_excel(result: dict, output_path: str):
    wb = _build_workbook(result)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def render_excel_bytes(result: dict) -> io.BytesIO:
    """Render analysis results as an Excel workbook in memory."""
    buf = io.BytesIO()
    wb = _build_workbook(result)
    wb.save(buf)
    buf.seek(0)
    return buf
